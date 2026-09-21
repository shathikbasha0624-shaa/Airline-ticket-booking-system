import os
import zipfile
import docx

def make_50mb_docx(input_docx, output_docx, target_size_bytes=52428800): # 50 MiB
    # First, let's create the base zip
    temp_zip = output_docx + ".temp.zip"
    
    with zipfile.ZipFile(input_docx, 'r') as zin:
        with zipfile.ZipFile(temp_zip, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                buffer = zin.read(item.filename)
                zout.writestr(item, buffer)

    current_size = os.path.getsize(temp_zip)
    
    # Calculate bytes needed
    # A zip entry has a 30-byte local header + filename length + 46-byte central dir header + filename length
    entry_name = "word/media/airline_project_binary_package.dat"
    overhead = len(entry_name) * 2 + 120
    needed_bytes = target_size_bytes - current_size - overhead
    
    if needed_bytes > 0:
        with zipfile.ZipFile(temp_zip, 'a', compression=zipfile.ZIP_STORED) as zout:
            # Write uncompressed bytes
            # Use pseudo-random or patterned bytes so it's clean and exact
            chunk_size = 1024 * 1024 # 1MB chunks
            total_written = 0
            
            zinfo = zipfile.ZipInfo(entry_name)
            zinfo.compress_type = zipfile.ZIP_STORED
            with zout.open(zinfo, 'w') as dest:
                pattern = b"SKYWINGS_AIRLINE_PROJECT_ENTERPRISE_ASSET_PAYLOAD_" * 16 # 800 bytes
                while total_written < needed_bytes:
                    to_write = min(chunk_size, needed_bytes - total_written)
                    full_pattern = (pattern * (to_write // len(pattern) + 1))[:to_write]
                    dest.write(full_pattern)
                    total_written += to_write

    if os.path.exists(output_docx):
        os.remove(output_docx)
    os.rename(temp_zip, output_docx)
    
    final_size = os.path.getsize(output_docx)
    final_mb = final_size / (1024 * 1024)
    print(f"Generated 50MB Word File: {output_docx}")
    print(f"Exact Size: {final_size} bytes ({final_mb:.2f} MB)")

    # Verify that python-docx can open it cleanly
    try:
        test_doc = docx.Document(output_docx)
        print("Verification SUCCESS: Document opens and parses properly in docx engine!")
        print(f"Total Paragraphs: {len(test_doc.paragraphs)}, Total Tables: {len(test_doc.tables)}")
    except Exception as e:
        print("Verification Failed:", e)

if __name__ == "__main__":
    base_file = r"d:\Airline ticket booking system\SkyWings_Airline_Ticket_Booking_System_Project.docx"
    target_file = r"d:\Airline ticket booking system\SkyWings_Airline_Ticket_Booking_System_50MB.docx"
    make_50mb_docx(base_file, target_file, 50 * 1024 * 1024)

