import os
import sys
import zipfile
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # -------------------------------------------------------------
    # COVER PAGE
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(80)
    title_p.paragraph_format.space_after = Pt(10)
    run_title = title_p.add_run("SKYWINGS AIRLINES")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(32)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(25)
    run_sub = sub_p.add_run("Full-Stack Enterprise Airline Ticket Booking & Management System")
    run_sub.font.size = Pt(16)
    run_sub.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.paragraph_format.space_before = Pt(30)
    meta_p.paragraph_format.space_after = Pt(120)
    run_meta = meta_p.add_run("Comprehensive Engineering Documentation & Source Code Specification\nPlatform: Node.js (Express) | Microsoft SQL Server | Modern Responsive Web")
    run_meta.font.size = Pt(12)
    run_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    author_p = doc.add_paragraph()
    author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_author = author_p.add_run("Project Implementation Report\nDatabase Server: DESKTOP-UVIULNA | Database: Airline\nVersion 2.0 (Full Release)")
    run_author.font.size = Pt(12)
    run_author.font.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # HELPER FUNCTIONS
    # -------------------------------------------------------------
    def add_heading_1(text):
        h = doc.add_heading(level=1)
        h.paragraph_format.space_before = Pt(20)
        h.paragraph_format.space_after = Pt(8)
        run = h.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)
        return h

    def add_heading_2(text):
        h = doc.add_heading(level=2)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run = h.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return h

    def add_heading_3(text):
        h = doc.add_heading(level=3)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        return h

    def add_code_block(title, content):
        add_heading_3(f"Source Code: {title}")
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, 'F1F5F9')
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(content)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # CHAPTER 1: EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    add_heading_1("1. Executive Summary & Project Overview")
    doc.add_paragraph(
        "SkyWings Airlines is an enterprise-grade flight booking, ticketing, and operations management application. "
        "Engineered with a client-server architecture, it seamlessly bridges an ultra-modern, responsive HTML5/CSS3/JavaScript "
        "frontend with a high-throughput Node.js Express backend, persisted onto an on-premises Microsoft SQL Server database."
    )
    doc.add_paragraph(
        "Key capabilities developed in this project include:\n"
        "• Dual User Roles: Passenger (Customer) and Administrator (System Operations).\n"
        "• Secure Credential Authentication: Salted hashing via Bcrypt and session persistence in LocalStorage.\n"
        "• Dynamic Multi-Tier Flight Booking: Economy, Business (2x fare), and First Class (3x fare) with live price calculations.\n"
        "• Instant Electronic Boarding Passes: Rendered digital boarding tickets with simulated barcode verification and print utility.\n"
        "• Operational Command Center: Administrative flight publishing, inventory control, and real-time customer reservation monitoring.\n"
        "• Windows Integrated Authentication: Native ODBC Driver 17 for SQL Server connectivity on DESKTOP-UVIULNA."
    )

    # -------------------------------------------------------------
    # CHAPTER 2: TECHNOLOGY STACK & ARCHITECTURE
    # -------------------------------------------------------------
    add_heading_1("2. System Architecture & Technical Specifications")
    doc.add_paragraph(
        "The system employs a decoupled, clean MVC-inspired architecture where server routes dispatch to dedicated controllers, "
        "connecting asynchronously to SQL Server through connection pools."
    )

    tbl = doc.add_table(rows=6, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer", "Technology", "Role & Key Features"]
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        set_cell_background(cell, '0F4C81')
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data = [
        ("Frontend Presentation", "HTML5, CSS3, Vanilla ES6+ JS", "Single Port Static Delivery, Glassmorphism, Micro-interactions"),
        ("Application Server", "Node.js (v22), Express (v5.x)", "RESTful API Endpoints, Routing Engine, Static Middleware"),
        ("Database Server", "Microsoft SQL Server 2017+", "Relational Schema, Transaction Integrity, Foreign Key Constraints"),
        ("Database Driver", "msnodesqlv8 (ODBC 17)", "Windows Native Authentication without plaintext credentials"),
        ("Security & Crypto", "Bcrypt.js, Dotenv", "Blowfish salted password hashing, Environment parameter isolation")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, col_val in enumerate(row_data):
            cell = tbl.cell(row_idx, col_idx)
            set_cell_background(cell, 'F8FAFC' if row_idx % 2 == 1 else 'FFFFFF')
            p = cell.paragraphs[0]
            p.add_run(col_val)

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # -------------------------------------------------------------
    # CHAPTER 3: DATABASE SCHEMA & DATA DICTIONARY
    # -------------------------------------------------------------
    add_heading_1("3. Database Schema & Data Dictionary")
    doc.add_paragraph("The relational database 'Airline' hosted on 'DESKTOP-UVIULNA' consists of three primary tables:")

    # Table 1: Users
    add_heading_2("3.1 Table: Users")
    tbl_users = doc.add_table(rows=6, cols=4)
    tbl_users.alignment = WD_TABLE_ALIGNMENT.CENTER
    u_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    for i, h in enumerate(u_headers):
        cell = tbl_users.cell(0, i)
        set_cell_background(cell, '1E3A8A')
        cell.paragraphs[0].add_run(h).font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    u_data = [
        ("Id", "INT IDENTITY(1,1)", "PRIMARY KEY", "Unique auto-incrementing user ID"),
        ("Username", "NVARCHAR(50)", "NOT NULL", "Passenger or administrator full name"),
        ("Email", "NVARCHAR(100)", "UNIQUE, NOT NULL", "Login email address"),
        ("PasswordHash", "NVARCHAR(255)", "NOT NULL", "Bcrypt salted hash string"),
        ("Role", "NVARCHAR(20)", "DEFAULT 'Customer'", "Authorization role ('Customer' or 'Admin')")
    ]
    for r_idx, row in enumerate(u_data, start=1):
        for c_idx, val in enumerate(row):
            cell = tbl_users.cell(r_idx, c_idx)
            set_cell_background(cell, 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF')
            cell.paragraphs[0].add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Table 2: Flights
    add_heading_2("3.2 Table: Flights")
    tbl_flights = doc.add_table(rows=7, cols=4)
    tbl_flights.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(u_headers):
        cell = tbl_flights.cell(0, i)
        set_cell_background(cell, '1E3A8A')
        cell.paragraphs[0].add_run(h).font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    f_data = [
        ("Id", "INT IDENTITY(1,1)", "PRIMARY KEY", "Unique flight record ID"),
        ("FlightNumber", "NVARCHAR(20)", "NOT NULL", "Commercial flight code (e.g. SW101)"),
        ("Origin", "NVARCHAR(50)", "NOT NULL", "Departure airport and city"),
        ("Destination", "NVARCHAR(50)", "NOT NULL", "Arrival airport and city"),
        ("DepartureTime", "DATETIME", "NOT NULL", "Scheduled departure timestamp"),
        ("Price", "DECIMAL(10,2)", "NOT NULL", "Base airfare in USD")
    ]
    for r_idx, row in enumerate(f_data, start=1):
        for c_idx, val in enumerate(row):
            cell = tbl_flights.cell(r_idx, c_idx)
            set_cell_background(cell, 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF')
            cell.paragraphs[0].add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Table 3: Bookings
    add_heading_2("3.3 Table: Bookings")
    tbl_bk = doc.add_table(rows=7, cols=4)
    tbl_bk.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(u_headers):
        cell = tbl_bk.cell(0, i)
        set_cell_background(cell, '1E3A8A')
        cell.paragraphs[0].add_run(h).font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    b_data = [
        ("Id", "INT IDENTITY(1,1)", "PRIMARY KEY", "Unique reservation number"),
        ("UserId", "INT", "FOREIGN KEY -> Users(Id)", "Reference to booking passenger"),
        ("FlightId", "INT", "FOREIGN KEY -> Flights(Id)", "Reference to reserved flight"),
        ("FlightClass", "NVARCHAR(50)", "DEFAULT 'Economy'", "Seat tier (Economy, Business, First Class)"),
        ("TotalPrice", "DECIMAL(10,2)", "DEFAULT 0.00", "Final billed amount including class multiplier & taxes"),
        ("Status", "NVARCHAR(20)", "DEFAULT 'Confirmed'", "Status ('Confirmed' or 'Cancelled')")
    ]
    for r_idx, row in enumerate(b_data, start=1):
        for c_idx, val in enumerate(row):
            cell = tbl_bk.cell(r_idx, c_idx)
            set_cell_background(cell, 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF')
            cell.paragraphs[0].add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # -------------------------------------------------------------
    # CHAPTER 4: REST API SPECIFICATION
    # -------------------------------------------------------------
    add_heading_1("4. REST API Endpoint Specifications")
    api_table = doc.add_table(rows=8, cols=4)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_headers = ["Method", "Route Endpoint", "Access", "Description"]
    for i, h in enumerate(api_headers):
        cell = api_table.cell(0, i)
        set_cell_background(cell, '0F4C81')
        cell.paragraphs[0].add_run(h).font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    api_data = [
        ("POST", "/api/auth/signup", "Public", "Registers new user with hashed password"),
        ("POST", "/api/auth/signin", "Public", "Authenticates user and returns user object + role"),
        ("GET", "/api/flights", "Public/User", "Retrieves all scheduled flights sorted by departure"),
        ("POST", "/api/flights", "Admin", "Publishes a new commercial flight to schedule"),
        ("DELETE", "/api/flights/:id", "Admin", "Deletes flight and associated bookings from database"),
        ("POST", "/api/bookings", "Customer", "Creates flight reservation with seat class and total fare"),
        ("GET", "/api/bookings/:userId", "Customer", "Retrieves booking history and e-tickets for active user")
    ]
    for r_idx, row in enumerate(api_data, start=1):
        for c_idx, val in enumerate(row):
            cell = api_table.cell(r_idx, c_idx)
            set_cell_background(cell, 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF')
            cell.paragraphs[0].add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # -------------------------------------------------------------
    # CHAPTER 5: SOURCE CODE LISTING
    # -------------------------------------------------------------
    add_heading_1("5. Complete Project Source Code Listings")
    doc.add_paragraph("The following sections contain the complete, verbatim source code for all project files.")

    project_root = r"d:\Airline ticket booking system"
    files_to_embed = [
        ("Server Configuration (.env)", os.path.join(project_root, "Server", ".env")),
        ("Database Setup Script (database_setup.sql)", os.path.join(project_root, "database_setup.sql")),
        ("Server Entry Point (server.js)", os.path.join(project_root, "Server", "server.js")),
        ("Database Connection Config (config/db.js)", os.path.join(project_root, "Server", "config", "db.js")),
        ("Authentication Controller (controllers/authController.js)", os.path.join(project_root, "Server", "controllers", "authController.js")),
        ("Flight Controller (controllers/flightController.js)", os.path.join(project_root, "Server", "controllers", "flightController.js")),
        ("Booking Controller (controllers/bookingController.js)", os.path.join(project_root, "Server", "controllers", "bookingController.js")),
        ("Auth Routes (routes/authRoutes.js)", os.path.join(project_root, "Server", "routes", "authRoutes.js")),
        ("Flight Routes (routes/flightRoutes.js)", os.path.join(project_root, "Server", "routes", "flightRoutes.js")),
        ("Booking Routes (routes/bookingRoutes.js)", os.path.join(project_root, "Server", "routes", "bookingRoutes.js")),
        ("Frontend Authentication UI (Client/index.html)", os.path.join(project_root, "Client", "index.html")),
        ("Frontend Authentication Logic (Client/app.js)", os.path.join(project_root, "Client", "app.js")),
        ("Passenger Dashboard UI (Client/dashboard.html)", os.path.join(project_root, "Client", "dashboard.html")),
        ("Passenger Dashboard Logic (Client/dashboard.js)", os.path.join(project_root, "Client", "dashboard.js")),
        ("Admin Control Center UI (Client/admin-dashboard.html)", os.path.join(project_root, "Client", "admin-dashboard.html")),
        ("Admin Control Center Logic (Client/admin-dashboard.js)", os.path.join(project_root, "Client", "admin-dashboard.js")),
        ("Design System & Stylesheet (Client/style.css)", os.path.join(project_root, "Client", "style.css")),
        ("Admin & Role Migration Script (Server/admin-migrate.js)", os.path.join(project_root, "Server", "admin-migrate.js")),
        ("Multi-Tier Schema Migration Script (Server/migrate-db.js)", os.path.join(project_root, "Server", "migrate-db.js"))
    ]

    for title, path in files_to_embed:
        if os.path.exists(path):
            try:
                with open(path, "r", encoding="utf-8") as f:
                    content = f.read()
                add_code_block(title, content)
            except Exception as e:
                add_code_block(title, f"Error reading file: {str(e)}")

    # -------------------------------------------------------------
    # CHAPTER 6: SETUP & VERIFICATION MANUAL
    # -------------------------------------------------------------
    add_heading_1("6. Installation, Configuration & Verification Guide")
    doc.add_paragraph(
        "1. Start SQL Server & verify Windows Authentication mode.\n"
        "2. Ensure ODBC Driver 17 for SQL Server is present in system ODBC Data Sources.\n"
        "3. Navigate to D:\\Airline ticket booking system\\Server and execute 'npm start'.\n"
        "4. Open web browser to http://localhost:3000 to launch the SkyWings Portal.\n"
        "5. Test Admin account with 'admin@gmail.com' and password '123456'.\n"
        "6. Test Customer account by registering or using demo chips."
    )

    doc_path = os.path.join(project_root, "SkyWings_Airline_Ticket_Booking_System_Project.docx")
    doc.save(doc_path)
    print(f"Base Word Document generated: {doc_path} ({os.path.getsize(doc_path)} bytes)")
    return doc_path

if __name__ == "__main__":
    create_document()

